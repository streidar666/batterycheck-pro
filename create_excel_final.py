from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# Configurazione
NUM_CELLE = 80
FILENAME = "Controllo_Codici_80.xlsx"

# Creazione del workbook
wb = Workbook()
ws = wb.active
ws.title = "Controllo Codici"

# Stili
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=12)
duplicate_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
neutral_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'), 
    right=Side(style='thin'), 
    top=Side(style='thin'), 
    bottom=Side(style='thin')
)
center_alignment = Alignment(horizontal='center', vertical='center')
left_alignment = Alignment(horizontal='left', vertical='center')

# Intestazioni
ws['A1'] = "N° Riga"
ws['B1'] = "CODICE DA INSERIRE"
ws['C1'] = "STATO"

# Applica stile intestazioni
for cell in ['A1', 'B1', 'C1']:
    ws[cell].fill = header_fill
    ws[cell].font = header_font
    ws[cell].alignment = center_alignment
    ws[cell].border = thin_border

# Generazione righe 1-80
for i in range(1, NUM_CELLE + 1):
    row = i + 1  # Excel row (start from 2)
    
    # Colonna A: Numerazione
    ws[f'A{row}'] = i
    ws[f'A{row}'].alignment = center_alignment
    ws[f'A{row}'].border = thin_border
    
    # Colonna B: Cella vuota per inserimento
    ws[f'B{row}'].border = thin_border
    
    # Colonna C: Formula per stato (vuoto se singolo, "DUPLICATO" se ripetuto)
    ws[f'C{row}'] = f'=IF(COUNTIF($B$2:$B${NUM_CELLE+1}, B{row})>1, "DUPLICATO", "")'
    ws[f'C{row}'].border = thin_border
    ws[f'C{row}'].alignment = center_alignment

# Formattazione Condizionale per evidenziare duplicati in rosso
# Nota: formula deve essere una lista/sequenza
duplicate_rule = FormulaRule(
    formula=["COUNTIF($B$2:$B$81, B2)>1"],
    fill=duplicate_fill
)
ws.conditional_formatting.add(f'B2:B{NUM_CELLE+1}', duplicate_rule)

# Impostazione larghezza colonne
ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 15

# --- SEZIONE RIEPILOGO ---
summary_start_row = NUM_CELLE + 5
ws[f'A{summary_start_row}'] = "RIEPILOGO CODICI"
ws[f'A{summary_start_row}'].font = Font(bold=True, size=14, color="4F81BD")
ws.merge_cells(f'A{summary_start_row}:B{summary_start_row}')

ws[f'A{summary_start_row+1}'] = "Codice"
ws[f'B{summary_start_row+1}'] = "Nr Totali"
ws[f'A{summary_start_row+1}'].font = Font(bold=True)
ws[f'B{summary_start_row+1}'].font = Font(bold=True)
ws[f'A{summary_start_row+1}'].border = thin_border
ws[f'B{summary_start_row+1}'].border = thin_border
ws[f'A{summary_start_row+1}'].alignment = center_alignment
ws[f'B{summary_start_row+1}'].alignment = center_alignment

# Nota per l'utente
note_row = summary_start_row - 2
ws[f'A{note_row}'] = "Nota: Le celle con codici duplicati si colorano automaticamente in ROSSO."
ws[f'A{note_row}'].font = Font(italic=True, color="9C0006", bold=True)

# Istruzioni
instr_row = summary_start_row - 1
ws[f'A{instr_row}'] = "Istruzioni: Inserisci i codici nella colonna B (righe 2-81)"
ws[f'A{instr_row}'].font = Font(italic=True)

# Salvataggio
wb.save(FILENAME)
print(f"✅ File '{FILENAME}' creato con successo!")
print("\n📋 Caratteristiche del file:")
print("   - Colonna A: Numerazione da 1 a 80")
print("   - Colonna B: Dove inserire i codici")
print("   - Colonna C: Indica 'DUPLICATO' se il codice è ripetuto")
print("   - Formattazione condizionale: Celle duplicate in ROSSO")
print("   - Sezione riepilogo pronta per i totali")
print("\n🔴 I codici uguali verranno evidenziati automaticamente in rosso!")
